"""
Excel Import Service
====================
Procesa archivos Excel subidos desde el dashboard y los vectoriza en Qdrant.
Soporta modo inmobiliario y moda/boutiques automáticamente.

Usado por: src/http/routers/excel.py → upload_excel()
"""

from __future__ import annotations

import io
import os
import uuid
import logging
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openai import OpenAI
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams, PointStruct

logger = logging.getLogger(__name__)

EMBEDDING_MODEL = "text-embedding-3-small"
VECTOR_DIMENSION = 1536
BATCH_SIZE = 100


def _get_qdrant_client() -> QdrantClient:
    url = os.getenv("QDRANT_URL")
    api_key = os.getenv("QDRANT_API_KEY")
    if not url or not api_key:
        raise ValueError("QDRANT_URL y QDRANT_API_KEY son requeridos")
    return QdrantClient(url=url, api_key=api_key)


def _get_openai_client() -> OpenAI:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("OPENAI_API_KEY es requerida")
    return OpenAI(api_key=api_key)


# ==================== PERSISTENCIA EN BD ====================

def _save_job_status(
    import_id: str,
    tenant_id: str,
    status: str,
    filename: str = "",
    total: int = 0,
    exitosos: int = 0,
    fallidos: int = 0,
    modo: str = "",
    error: str = "",
) -> None:
    """Guarda o actualiza el estado del job en excel_import_jobs."""
    try:
        from src.core.database import get_db
        from datetime import datetime, timezone

        db = get_db().client
        now = datetime.now(timezone.utc).isoformat()

        existing = (
            db.table("excel_import_jobs")
            .select("id")
            .eq("id", import_id)
            .execute()
        )

        data = {
            "status": status,
            "total": total,
            "exitosos": exitosos,
            "fallidos": fallidos,
            "modo": modo,
            "error": error[:500] if error else None,
            "updated_at": now,
        }

        if existing.data:
            db.table("excel_import_jobs").update(data).eq("id", import_id).execute()
        else:
            data.update({
                "id": import_id,
                "tenant_id": tenant_id,
                "filename": filename,
                "created_at": now,
            })
            db.table("excel_import_jobs").insert(data).execute()

    except Exception as e:
        logger.error(f"[{import_id}] Error guardando job status: {e}")


def get_job_status(import_id: str, tenant_id: str) -> Optional[Dict[str, Any]]:
    """
    Obtiene el estado real de un job desde excel_import_jobs.
    Usado por GET /plugins/excel/imports/{import_id}
    """
    try:
        from src.core.database import get_db
        db = get_db().client

        result = (
            db.table("excel_import_jobs")
            .select("*")
            .eq("id", import_id)
            .eq("tenant_id", tenant_id)
            .limit(1)
            .execute()
        )

        return result.data[0] if result.data else None

    except Exception as e:
        logger.error(f"Error obteniendo job status {import_id}: {e}")
        return None


def get_jobs_by_tenant(
    tenant_id: str,
    limit: int = 10,
    offset: int = 0,
    status_filter: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """
    Lista jobs de importación de un tenant.
    Usado por GET /plugins/excel/imports
    """
    try:
        from src.core.database import get_db
        db = get_db().client

        query = (
            db.table("excel_import_jobs")
            .select("*")
            .eq("tenant_id", tenant_id)
            .order("created_at", desc=True)
            .limit(limit)
            .offset(offset)
        )

        if status_filter:
            query = query.eq("status", status_filter)

        result = query.execute()
        return result.data or []

    except Exception as e:
        logger.error(f"Error listando jobs para tenant {tenant_id}: {e}")
        return []


# ==================== DETECCIÓN DE MODO ====================

def _detectar_modo(headers: List[str]) -> str:
    columnas_moda = {"tallas", "colores", "referencia", "categoria", "talla", "color"}
    columnas_inmobiliario = {"barrio", "habitaciones", "costo del arriendo", "bano"}
    headers_set = set(headers)
    if columnas_moda.intersection(headers_set):
        return "moda"
    if columnas_inmobiliario.intersection(headers_set):
        return "inmobiliario"
    return "moda"


# ==================== BUILD TEXT ====================

def _build_text_inmobiliario(p: dict) -> str:
    disponible_raw = p.get("disponible", p.get("stock", 1))
    disponible = "Disponible" if str(disponible_raw).lower() in ("1", "si", "sí", "yes", "true") else "No disponible"
    partes = [
        f"Propiedad: {p.get('name') or p.get('nombre', '')}",
        f"Tipo: {p.get('tipo de inmueble') or p.get('tipo', '')}",
        f"Barrio: {p.get('barrio', '')}",
        f"Arriendo: ${p.get('costo del arriendo') or p.get('precio', 0):,} COP",
        f"Habitaciones: {p.get('habitaciones', 0)}",
        f"Baños: {p.get('bano') or p.get('banos', 0)}",
        f"Parqueadero: {p.get('parqueadero', '')}",
        f"Mascotas: {p.get('mascotas', '')}",
        f"Estado: {disponible}",
    ]
    desc = p.get("description") or p.get("descripcion", "")
    if desc:
        partes.append(f"Descripción: {desc}")
    return "\n".join(partes).strip()


def _build_text_moda(p: dict) -> str:
    disponible_raw = p.get("disponible", "si")
    disponible = "Disponible" if str(disponible_raw).lower() in ("si", "sí", "yes", "1", "true") else "No disponible"
    precio = p.get("precio", 0) or 0
    descuento = p.get("descuento_pct", 0) or 0
    partes = [
        f"Producto: {p.get('nombre') or p.get('name', '')}",
        f"Referencia: {p.get('referencia', '')}",
        f"Categoría: {p.get('categoria') or p.get('category', '')}",
        f"Precio: ${int(precio):,} COP",
    ]
    if descuento and float(descuento) > 0:
        precio_final = int(precio) * (1 - float(descuento) / 100)
        partes.append(f"Descuento: {descuento}% — Precio con descuento: ${precio_final:,.0f} COP")
    tallas = p.get("tallas") or p.get("talla", "")
    colores = p.get("colores") or p.get("color", "")
    if tallas:
        partes.append(f"Tallas disponibles: {tallas}")
    if colores:
        partes.append(f"Colores disponibles: {colores}")
    partes.append(f"Stock: {int(p.get('stock', 0) or 0)} unidades")
    partes.append(f"Estado: {disponible}")
    desc = p.get("descripcion") or p.get("description", "")
    if desc:
        partes.append(f"Descripción: {desc}")
    return "\n".join(partes).strip()


def _build_metadata_inmobiliario(p: dict, tenant_id: str) -> dict:
    disponible_raw = p.get("disponible", p.get("stock", 1))
    disponible = str(disponible_raw).lower() in ("1", "si", "sí", "yes", "true")
    return {
        "name": p.get("name") or p.get("nombre", ""),
        "tipo": p.get("tipo de inmueble") or p.get("tipo", ""),
        "barrio": p.get("barrio", ""),
        "precio": p.get("costo del arriendo") or p.get("precio", 0),
        "habitaciones": p.get("habitaciones", 0),
        "banos": p.get("bano") or p.get("banos", 0),
        "parqueadero": p.get("parqueadero", ""),
        "mascotas": p.get("mascotas", ""),
        "disponible": disponible,
        "description": p.get("description") or p.get("descripcion", ""),
        "tenant_id": tenant_id,
        "modo": "inmobiliario",
    }


def _build_metadata_moda(p: dict, tenant_id: str) -> dict:
    disponible_raw = p.get("disponible", "si")
    disponible = str(disponible_raw).lower() in ("si", "sí", "yes", "1", "true")
    precio = p.get("precio", 0) or 0
    descuento = p.get("descuento_pct", 0) or 0
    return {
        "name": p.get("nombre") or p.get("name", ""),
        "referencia": p.get("referencia", ""),
        "categoria": p.get("categoria") or p.get("category", ""),
        "precio": int(precio),
        "tallas": p.get("tallas") or p.get("talla", ""),
        "colores": p.get("colores") or p.get("color", ""),
        "stock": int(p.get("stock", 0) or 0),
        "disponible": disponible,
        "descuento_pct": float(descuento),
        "description": p.get("descripcion") or p.get("description", ""),
        "tenant_id": tenant_id,
        "modo": "moda",
    }


def _leer_excel(contents: bytes) -> Tuple[List[str], List[dict]]:
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True)
    if "productos" in wb.sheetnames:
        ws = wb["productos"]
    else:
        ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    items = []
    for row in rows[1:]:
        d = dict(zip(headers, row))
        nombre = d.get("nombre") or d.get("name") or d.get("nombre_producto")
        if nombre:
            items.append(d)
    return headers, items


# ==================== PROCESO PRINCIPAL ====================

def process_excel_upload(
    contents: bytes,
    tenant_id: str,
    tenant_slug: str = "",
    import_id: str = "",
    filename: str = "",
) -> Dict[str, Any]:
    """
    Procesa un archivo Excel subido desde el dashboard.
    Persiste el estado en excel_import_jobs para polling real.
    """
    collection_name = f"products_{tenant_id}"
    errores = []
    exitosos = 0

    # ✅ Guardar job como processing al inicio
    _save_job_status(
        import_id=import_id,
        tenant_id=tenant_id,
        status="processing",
        filename=filename,
    )

    try:
        headers, items = _leer_excel(contents)

        if not items:
            _save_job_status(import_id=import_id, tenant_id=tenant_id, status="failed",
                             filename=filename, error="El archivo Excel está vacío")
            return {"import_id": import_id, "status": "failed", "total": 0,
                    "exitosos": 0, "fallidos": 0, "modo": "desconocido",
                    "error": "El archivo Excel está vacío"}

        modo = _detectar_modo(headers)
        logger.info(f"[{import_id}] Modo: {modo} — {len(items)} items")

        documents = []
        for i, item in enumerate(items):
            try:
                nombre = item.get("nombre") or item.get("name") or f"item_{i}"
                doc_id = f"{tenant_slug}_{nombre}_{i}".replace(" ", "_").lower()
                if modo == "moda":
                    text = _build_text_moda(item)
                    metadata = _build_metadata_moda(item, tenant_id)
                else:
                    text = _build_text_inmobiliario(item)
                    metadata = _build_metadata_inmobiliario(item, tenant_id)
                documents.append({"id": doc_id, "text": text, "metadata": metadata})
            except Exception as e:
                errores.append({"fila": i + 2, "error": str(e)})

        if not documents:
            _save_job_status(import_id=import_id, tenant_id=tenant_id, status="failed",
                             filename=filename, total=len(items), fallidos=len(items), modo=modo,
                             error="No se pudieron procesar los items")
            return {"import_id": import_id, "status": "failed", "total": len(items),
                    "exitosos": 0, "fallidos": len(items), "modo": modo}

        # Embeddings
        openai_client = _get_openai_client()
        texts = [doc["text"] for doc in documents]
        vectors = []
        for i in range(0, len(texts), BATCH_SIZE):
            batch = texts[i:i + BATCH_SIZE]
            response = openai_client.embeddings.create(input=batch, model=EMBEDDING_MODEL)
            vectors.extend([item.embedding for item in response.data])

        # Qdrant
        qdrant = _get_qdrant_client()
        collections = [c.name for c in qdrant.get_collections().collections]
        if collection_name in collections:
            qdrant.delete_collection(collection_name=collection_name)
        qdrant.create_collection(
            collection_name=collection_name,
            vectors_config=VectorParams(size=VECTOR_DIMENSION, distance=Distance.COSINE),
        )

        points = []
        for doc, vector in zip(documents, vectors):
            point_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, doc["id"]))
            payload = {"custom_id": doc["id"], "text": doc["text"], **doc["metadata"]}
            points.append(PointStruct(id=point_id, vector=vector, payload=payload))
        qdrant.upsert(collection_name=collection_name, points=points)
        exitosos = len(points)

        # ✅ Marcar como completed
        _save_job_status(
            import_id=import_id, tenant_id=tenant_id, status="completed",
            filename=filename, total=len(items), exitosos=exitosos,
            fallidos=len(errores), modo=modo,
        )

        return {
            "import_id": import_id, "status": "completed", "total": len(items),
            "exitosos": exitosos, "fallidos": len(errores), "modo": modo,
            "collection": collection_name, "errores": errores,
        }

    except Exception as e:
        logger.error(f"[{import_id}] Error crítico: {e}")
        # ✅ Marcar como failed
        _save_job_status(import_id=import_id, tenant_id=tenant_id, status="failed",
                         filename=filename, error=str(e))
        return {"import_id": import_id, "status": "failed", "total": 0,
                "exitosos": 0, "fallidos": 0, "modo": "desconocido", "error": str(e)}